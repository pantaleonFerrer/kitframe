"""Fill FC Barcelona 'Descripcions producte per ecom' workbooks.

Green / non-red rows missing Spanish description are filled. Red rows are skipped.
Locales in the sheet: ES, CA, EN, FR, DE, JP, KR (+ meta titles/descriptions).

Descriptions are grounded in product photos; fit / measures / care come from
operator notes when provided (translated to all locales).
"""
from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass
from pathlib import PurePosixPath
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.workbook.workbook import Workbook

from kitframe.extraction.schema import enforce_strict_schema

ECOM_LOCALES = ("ES", "CA", "EN", "FR", "DE", "JP", "KR")

LANG_BLOCKS = {
    "ES": dict(title="T", desc="U", fit="V", measures="W", care="X"),
    "CA": dict(title="Z", desc="AA", fit="AB", measures="AC", care="AD"),
    "EN": dict(title="AF", desc="AG", fit="AH", measures="AI", care="AJ"),
    "FR": dict(title="AL", desc="AM", fit="AN", measures="AO", care="AP"),
    "DE": dict(title="AR", desc="AS", fit="AT", measures="AU", care="AV"),
    "JP": dict(title="AX", desc="AY", fit="AZ", measures="BA", care="BB"),
    "KR": dict(title="BD", desc="BE", fit="BF", measures="BG", care="BH"),
}

META = {
    "CA": ("BJ", "BQ"),
    "ES": ("BK", "BR"),
    "EN": ("BL", "BS"),
    "DE": ("BM", "BT"),
    "FR": ("BN", "BU"),
    "JA": ("BO", "BV"),
    "KO": ("BP", "BW"),
}

META_SRC = {"JA": "JP", "KO": "KR"}

DEFAULT_FIT = {
    "ES": "Regular fit: elige tu talla habitual.",
    "CA": "Regular fit: escull la teva talla habitual.",
    "EN": "Regular fit: choose your usual size.",
    "FR": "Coupe regular : choisis ta taille habituelle.",
    "DE": "Regular Fit: wähle deine übliche Größe.",
    "JP": "レギュラーフィット：いつものサイズを選んでください。",
    "KR": "레귤러 핏: 평소 사이즈를 선택하세요.",
}

IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
MAX_IMAGES_PER_REF = 8


@dataclass
class EcomRow:
    row: int
    ref: str
    title: str
    collection: str
    product_type: str
    existing_fit_es: str


def _cell(ws: Any, row: int, letter: str) -> Any:
    return ws.cell(row, column_index_from_string(letter))


def _is_red(cell: Any) -> bool:
    fill = cell.fill
    if not fill or fill.patternType is None or not fill.fgColor:
        return False
    rgb = getattr(fill.fgColor, "rgb", None)
    return bool(rgb and str(rgb).upper().endswith("FF0000"))


def load_workbook_bytes(data: bytes) -> Workbook:
    return load_workbook(io.BytesIO(data))


def workbook_to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def list_rows_to_fill(wb: Workbook, *, only_missing_es_desc: bool = True) -> list[EcomRow]:
    ws = wb.active
    out: list[EcomRow] = []
    for row in range(4, (ws.max_row or 0) + 1):
        ref_cell = ws.cell(row, 1)
        ref = ref_cell.value
        if ref is None or str(ref).strip() == "":
            continue
        if _is_red(ref_cell):
            continue
        es_desc = _cell(ws, row, "U").value
        if only_missing_es_desc and es_desc and str(es_desc).strip():
            continue
        out.append(
            EcomRow(
                row=row,
                ref=str(ref).strip(),
                title=str(ws.cell(row, 3).value or "").strip(),
                collection=str(ws.cell(row, 4).value or "").strip(),
                product_type=str(ws.cell(row, 5).value or "").strip(),
                existing_fit_es=str(_cell(ws, row, "V").value or "").strip(),
            )
        )
    return out


def normalize_ref_key(value: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", value.strip().upper())


def is_product_image_filename(filename: str) -> bool:
    suffix = PurePosixPath(filename).suffix.lower()
    return suffix in IMAGE_SUFFIXES


def match_filename_to_ref(filename: str, refs: list[str]) -> str | None:
    """Map `BLMP813081_front.jpg` / `blmp813081-1.png` to a pending ref."""
    if not refs or not is_product_image_filename(filename):
        return None
    stem = PurePosixPath(filename).stem
    # Longer refs first so BLMP813081A beats BLMP813081
    ordered = sorted(refs, key=len, reverse=True)
    for ref in ordered:
        if not ref.strip():
            continue
        pattern = re.compile(
            r"^" + re.escape(ref.strip()) + r"(?:$|[_\-\s.].*)",
            re.IGNORECASE,
        )
        if pattern.match(stem):
            return ref
    return None


def group_images_by_ref(
    filenames: list[str],
    refs: list[str],
) -> tuple[dict[str, list[str]], list[str]]:
    """Return (ref -> filenames, unmatched filenames)."""
    grouped: dict[str, list[str]] = {r: [] for r in refs}
    unmatched: list[str] = []
    for name in filenames:
        ref = match_filename_to_ref(name, refs)
        if ref is None:
            unmatched.append(name)
            continue
        if len(grouped[ref]) < MAX_IMAGES_PER_REF:
            grouped[ref].append(name)
    return grouped, unmatched


def parse_notes_json(raw: str | None) -> dict[str, dict[str, str]]:
    """Parse notes keyed by product ref: {fit, care, measures, extra}."""
    if not raw or not str(raw).strip():
        return {}
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ValueError(f"notes_json is not valid JSON: {exc}") from exc
    if not isinstance(data, dict):
        raise ValueError("notes_json must be an object keyed by product reference")
    out: dict[str, dict[str, str]] = {}
    for key, value in data.items():
        ref = str(key).strip()
        if not ref:
            continue
        if isinstance(value, str):
            out[ref] = {"extra": value.strip()}
            continue
        if not isinstance(value, dict):
            continue
        block: dict[str, str] = {}
        for field in ("fit", "care", "measures", "extra"):
            v = value.get(field)
            if v is not None and str(v).strip():
                block[field] = str(v).strip()
        if block:
            out[ref] = block
    return out


def ecom_row_schema() -> dict[str, Any]:
    locale_block = {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "title": {"type": "string"},
            "description": {"type": "string"},
            "fit": {"type": "string"},
            "measures": {"type": "string"},
            "care": {"type": "string"},
            "meta_title": {"type": "string"},
            "meta_description": {"type": "string"},
        },
        "required": [
            "title",
            "description",
            "fit",
            "measures",
            "care",
            "meta_title",
            "meta_description",
        ],
    }
    return enforce_strict_schema(
        {
            "type": "object",
            "additionalProperties": False,
            "properties": {loc: locale_block for loc in ECOM_LOCALES},
            "required": list(ECOM_LOCALES),
        }
    )


def apply_row_fill(wb: Workbook, row: int, content: dict[str, Any]) -> None:
    ws = wb.active
    for lang, cols in LANG_BLOCKS.items():
        block = content.get(lang) or {}
        if not isinstance(block, dict):
            continue
        title = str(block.get("title") or "").strip()
        desc = str(block.get("description") or "").strip()
        fit = str(block.get("fit") or "").strip() or DEFAULT_FIT.get(lang, "")
        measures = str(block.get("measures") or "").strip()
        care = str(block.get("care") or "").strip()
        if title:
            _cell(ws, row, cols["title"]).value = title
        if desc:
            _cell(ws, row, cols["desc"]).value = desc
        if fit:
            _cell(ws, row, cols["fit"]).value = fit
        if measures:
            _cell(ws, row, cols["measures"]).value = measures
        if care:
            _cell(ws, row, cols["care"]).value = care

    for meta_lang, (title_col, desc_col) in META.items():
        src = META_SRC.get(meta_lang, meta_lang)
        block = content.get(src) or {}
        if not isinstance(block, dict):
            continue
        meta_title = str(block.get("meta_title") or block.get("title") or "").strip()
        meta_desc = str(block.get("meta_description") or "").strip()
        if meta_title:
            _cell(ws, row, title_col).value = meta_title
        if meta_desc:
            _cell(ws, row, desc_col).value = meta_desc


def build_system_prompt(*, tone: str, do_not_translate: list[str]) -> str:
    return (
        "You write ecommerce product copy for FC Barcelona merchandising "
        "(official Club / co-branded / licensed products). "
        f"Tone: {tone or 'clear, commercial, factual'}. "
        f"Never alter or invent translations for these brand terms: "
        f"{do_not_translate or ['FC Barcelona', 'Barça', 'Blaugrana']}. "
        "CRITICAL: Product descriptions MUST be grounded in the attached product photos — "
        "describe colour, graphics, silhouette, collar, print and visible construction you see. "
        "Derive a natural commercial product name from the technical title (column C). "
        "For fit / measures / care: when the operator provides notes, translate those facts "
        "faithfully into every locale and do NOT invent numbers or care symbols. "
        "If notes omit a field, leave measures empty and use cautious care wording "
        "(follow the care label) rather than inventing fabric percentages. "
        "Accessories may use N/A for fit when appropriate. "
        "meta_title <= 70 chars, meta_description <= 160 chars. "
        "Return ALL locales: ES, CA, EN, FR, DE, JP, KR. "
        "Plain text only (no HTML)."
    )


def build_user_prompt(
    row: EcomRow,
    *,
    notes: dict[str, str] | None = None,
    image_count: int = 0,
) -> str:
    notes = notes or {}
    payload = {
        "referencia": row.ref,
        "titulo_tecnico": row.title,
        "coleccion": row.collection,
        "tipo_producto": row.product_type,
        "fit_es_existente": row.existing_fit_es or None,
        "fotos_adjuntas": image_count,
        "notas_operador": {
            "fit": notes.get("fit") or None,
            "medidas": notes.get("measures") or None,
            "cuidados": notes.get("care") or None,
            "extra": notes.get("extra") or None,
        },
        "instrucciones": (
            "Write copy from the photos. Use notas_operador for fit/measures/care when present. "
            "Translate notes into every locale."
        ),
    }
    return json.dumps(payload, ensure_ascii=False)
