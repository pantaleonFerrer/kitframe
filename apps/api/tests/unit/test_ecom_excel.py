from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from kitframe.content.ecom_excel import (
    apply_row_fill,
    build_user_prompt,
    group_images_by_ref,
    list_rows_to_fill,
    load_workbook_bytes,
    match_filename_to_ref,
    parse_notes_json,
    workbook_to_bytes,
)


def _sample_wb() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws["A3"] = "Referencia"
    ws["C3"] = "Título"
    ws["T3"] = "Título ES"
    ws["U3"] = "Desccripción ES"
    # green-ish row to fill
    ws["A4"] = "SKU-GREEN"
    ws["C4"] = "TEE WALL TEST"
    ws["D4"] = "TEE WALL"
    ws["E4"] = "Propis"
    # red row (skip)
    ws["A5"] = "SKU-RED"
    ws["C5"] = "ALREADY DONE"
    ws["A5"].fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
    ws["U5"] = "Already filled"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_list_rows_skips_red_and_filled() -> None:
    wb = load_workbook_bytes(_sample_wb())
    rows = list_rows_to_fill(wb)
    assert [r.ref for r in rows] == ["SKU-GREEN"]


def test_apply_row_fill_writes_es() -> None:
    wb = load_workbook_bytes(_sample_wb())
    apply_row_fill(
        wb,
        4,
        {
            "ES": {
                "title": "Camiseta test",
                "description": "Descripción de prueba",
                "fit": "Regular fit",
                "measures": "",
                "care": "Algodón",
                "meta_title": "Meta",
                "meta_description": "Meta desc",
            }
        },
    )
    data = workbook_to_bytes(wb)
    wb2 = load_workbook_bytes(data)
    assert wb2.active["T4"].value == "Camiseta test"
    assert wb2.active["U4"].value == "Descripción de prueba"
    assert wb2.active["BK4"].value == "Meta"
    assert wb2.active["BR4"].value == "Meta desc"


def test_match_filename_to_ref() -> None:
    refs = ["BLMP813081", "BLMP813082"]
    assert match_filename_to_ref("BLMP813081.jpg", refs) == "BLMP813081"
    assert match_filename_to_ref("blmp813081_1.png", refs) == "BLMP813081"
    assert match_filename_to_ref("BLMP813082-front.webp", refs) == "BLMP813082"
    assert match_filename_to_ref("random.jpg", refs) is None
    assert match_filename_to_ref("BLMP81308.jpg", refs) is None  # prefix too short


def test_group_images_by_ref() -> None:
    refs = ["BLMP813081", "BLMP813082"]
    grouped, unmatched = group_images_by_ref(
        ["BLMP813081.jpg", "BLMP813081_2.png", "other.gif", "BLMP813082.jpg"],
        refs,
    )
    assert [f for f in grouped["BLMP813081"]] == ["BLMP813081.jpg", "BLMP813081_2.png"]
    assert grouped["BLMP813082"] == ["BLMP813082.jpg"]
    assert unmatched == ["other.gif"]


def test_parse_notes_json() -> None:
    notes = parse_notes_json(
        '{"BLMP813081": {"fit": "Regular", "care": "30C", "extra": "official"}}'
    )
    assert notes["BLMP813081"]["fit"] == "Regular"
    assert notes["BLMP813081"]["care"] == "30C"
    assert parse_notes_json('{"X": "free text"}')["X"]["extra"] == "free text"


def test_build_user_prompt_includes_notes() -> None:
    wb = load_workbook_bytes(_sample_wb())
    row = list_rows_to_fill(wb)[0]
    user = build_user_prompt(
        row,
        notes={"fit": "Slim", "care": "Wash 30C", "measures": "Chest 52"},
        image_count=2,
    )
    assert "SKU-GREEN" in user
    assert "Slim" in user
    assert '"fotos_adjuntas": 2' in user
