#!/usr/bin/env python3
"""Fill green (to-do) rows in the FC Barcelona 'Descripcions producte per ecom' workbook.

Skips red rows (already done / out of scope). Writes a new xlsx next to the input
(or --output). Usage:

  python scripts/fill_descripcions_ecom.py "/path/to/Descripcions producte per ecom.xlsx"
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

# Ecommerce copy columns (header row 3)
LANG_BLOCKS = {
    "ES": dict(title="T", desc="U", fit="V", measures="W", care="X"),
    "CA": dict(title="Z", desc="AA", fit="AB", measures="AC", care="AD"),
    "EN": dict(title="AF", desc="AG", fit="AH", measures="AI", care="AJ"),
    "FR": dict(title="AL", desc="AM", fit="AN", measures="AO", care="AP"),
    "DE": dict(title="AR", desc="AS", fit="AT", measures="AU", care="AV"),
    "JP": dict(title="AX", desc="AY", fit="AZ", measures="BA", care="BB"),
    "KR": dict(title="BD", desc="BE", fit="BF", measures="BG", care="BH"),
}

META = {
    "CA": ("BJ", "BQ"),
    "ES": ("BK", "BR"),
    "EN": ("BL", "BS"),
    "DE": ("BM", "BT"),
    "FR": ("BN", "BU"),
    "JA": ("BO", "BV"),
    "KO": ("BP", "BW"),
}

FIT = {
    "ES": "Regular fit: elige tu talla habitual.",
    "CA": "Regular fit: escull la teva talla habitual.",
    "EN": "Regular fit: choose your usual size.",
    "FR": "Coupe regular : choisis ta taille habituelle.",
    "DE": "Regular Fit: wähle deine übliche Größe.",
    "JP": "レギュラーフィット：いつものサイズを選んでください。",
    "KR": "레귤러 핏: 평소 사이즈를 선택하세요.",
}

CARE_TEE = {
    "ES": "Tejido de algodón. Seguir las instrucciones de la etiqueta de cuidados.",
    "CA": "Teixit de cotó. Seguir les instruccions de l’etiqueta de cures.",
    "EN": "Cotton fabric. Follow the care label instructions.",
    "FR": "Tissu en coton. Suivre les instructions de l’étiquette d’entretien.",
    "DE": "Baumwollstoff. Pflegehinweise auf dem Etikett beachten.",
    "JP": "コットン素材。ケアラベルの指示に従ってください。",
    "KR": "코튼 소재. 케어 라벨의 지침을 따르세요.",
}

MEASURES_NA = {
    "ES": "",
    "CA": "",
    "EN": "",
    "FR": "",
    "DE": "",
    "JP": "",
    "KR": "",
}

# Per-SKU commercial copy. Keys = Referencia.
PRODUCTS: dict[str, dict[str, dict[str, str]]] = {
    "BLMP813081": {
        "ES": {
            "title": "Camiseta FC Barcelona Tee Wall III Match Beige",
            "desc": (
                "Camiseta de manga corta de la colección Tee Wall III en tono beige. "
                "Diseño gráfico Match en el delantero y marca FC Barcelona. "
                "Corte regular para un uso diario cómodo."
            ),
            "meta_desc": (
                "Camiseta beige FC Barcelona Tee Wall III Match de manga corta, "
                "corte regular y gráfico frontal."
            ),
        },
        "CA": {
            "title": "Samarreta FC Barcelona Tee Wall III Match Beige",
            "desc": (
                "Samarreta de màniga curta de la col·lecció Tee Wall III en to beige. "
                "Disseny gràfic Match al davant i marca FC Barcelona. "
                "Tall regular per a un ús diari còmode."
            ),
            "meta_desc": (
                "Samarreta beige FC Barcelona Tee Wall III Match de màniga curta, "
                "tall regular i gràfic frontal."
            ),
        },
        "EN": {
            "title": "FC Barcelona Tee Wall III Match Beige T-Shirt",
            "desc": (
                "Short-sleeve t-shirt from the Tee Wall III collection in a beige tone. "
                "Features a Match graphic on the front and FC Barcelona branding. "
                "Regular fit for everyday wear."
            ),
            "meta_desc": (
                "Beige FC Barcelona Tee Wall III Match short-sleeve t-shirt with "
                "regular fit and front graphic."
            ),
        },
        "FR": {
            "title": "T-shirt FC Barcelona Tee Wall III Match Beige",
            "desc": (
                "T-shirt à manches courtes de la collection Tee Wall III en beige. "
                "Motif graphique Match sur le devant et branding FC Barcelona. "
                "Coupe regular pour un port quotidien confortable."
            ),
            "meta_desc": (
                "T-shirt beige FC Barcelona Tee Wall III Match à manches courtes, "
                "coupe regular et graphique frontal."
            ),
        },
        "DE": {
            "title": "FC Barcelona Tee Wall III Match Beige T-Shirt",
            "desc": (
                "Kurzarm-T-Shirt der Tee Wall III Kollektion in Beige. "
                "Mit Match-Grafik auf der Vorderseite und FC Barcelona Branding. "
                "Regular Fit für den Alltag."
            ),
            "meta_desc": (
                "Beiges FC Barcelona Tee Wall III Match Kurzarm-T-Shirt mit "
                "Regular Fit und Frontgrafik."
            ),
        },
        "JP": {
            "title": "FCバルセロナ Tee Wall III Match ベージュ Tシャツ",
            "desc": (
                "Tee Wall IIIコレクションのベージュの半袖Tシャツ。"
                "フロントにMatchグラフィックとFCバルセロナのブランディングを配置。"
                "デイリーユースに適したレギュラーフィット。"
            ),
            "meta_desc": (
                "ベージュのFCバルセロナ Tee Wall III Match半袖Tシャツ。"
                "レギュラーフィットとフロントグラフィック。"
            ),
        },
        "KR": {
            "title": "FC 바르셀로나 Tee Wall III Match 베이지 티셔츠",
            "desc": (
                "Tee Wall III 컬렉션의 베이지 반팔 티셔츠입니다. "
                "전면 Match 그래픽과 FC 바르셀로나 브랜딩이 적용되었습니다. "
                "일상 착용에 적합한 레귤러 핏입니다."
            ),
            "meta_desc": (
                "베이지 FC 바르셀로나 Tee Wall III Match 반팔 티셔츠. "
                "레귤러 핏과 전면 그래픽."
            ),
        },
    },
    "BLMP813082": {
        "ES": {
            "title": "Polo FC Barcelona Tee Wall III Blaugrana",
            "desc": (
                "Polo de manga corta de la colección Tee Wall III en color blaugrana. "
                "Cuello tipo polo, acabado limpio y marca FC Barcelona. "
                "Ideal para un look casual con identidad del Club."
            ),
            "meta_desc": (
                "Polo blaugrana FC Barcelona Tee Wall III de manga corta con "
                "cuello polo y marca del Club."
            ),
        },
        "CA": {
            "title": "Polo FC Barcelona Tee Wall III Blaugrana",
            "desc": (
                "Polo de màniga curta de la col·lecció Tee Wall III en color blaugrana. "
                "Coll tipus polo, acabat net i marca FC Barcelona. "
                "Ideal per a un look casual amb identitat del Club."
            ),
            "meta_desc": (
                "Polo blaugrana FC Barcelona Tee Wall III de màniga curta amb "
                "coll polo i marca del Club."
            ),
        },
        "EN": {
            "title": "FC Barcelona Tee Wall III Blaugrana Polo",
            "desc": (
                "Short-sleeve polo from the Tee Wall III collection in blaugrana. "
                "Classic polo collar, clean finish and FC Barcelona branding. "
                "A casual look with Club identity."
            ),
            "meta_desc": (
                "Blaugrana FC Barcelona Tee Wall III short-sleeve polo with "
                "polo collar and Club branding."
            ),
        },
        "FR": {
            "title": "Polo FC Barcelona Tee Wall III Blaugrana",
            "desc": (
                "Polo à manches courtes de la collection Tee Wall III en blaugrana. "
                "Col polo classique, finition nette et branding FC Barcelona. "
                "Look casual à l’identité du Club."
            ),
            "meta_desc": (
                "Polo blaugrana FC Barcelona Tee Wall III à manches courtes avec "
                "col polo et branding du Club."
            ),
        },
        "DE": {
            "title": "FC Barcelona Tee Wall III Blaugrana Polo",
            "desc": (
                "Kurzarm-Polo der Tee Wall III Kollektion in Blaugrana. "
                "Klassischer Polokragen, cleanes Finish und FC Barcelona Branding. "
                "Casual Look mit Club-Identität."
            ),
            "meta_desc": (
                "Blaugrana FC Barcelona Tee Wall III Kurzarm-Polo mit "
                "Polokragen und Club-Branding."
            ),
        },
        "JP": {
            "title": "FCバルセロナ Tee Wall III Blaugrana ポロ",
            "desc": (
                "Tee Wall IIIコレクションのブラウグラナ半袖ポロ。"
                "クラシックなポロカラー、クリーンな仕上げ、FCバルセロナのブランディング。"
                "クラブのアイデンティティを感じるカジュアルルックに。"
            ),
            "meta_desc": (
                "ブラウグラナのFCバルセロナ Tee Wall III半袖ポロ。"
                "ポロカラーとクラブブランディング。"
            ),
        },
        "KR": {
            "title": "FC 바르셀로나 Tee Wall III Blaugrana 폴로",
            "desc": (
                "Tee Wall III 컬렉션의 블라우그라나 반팔 폴로입니다. "
                "클래식 폴로 칼라, 클린한 마감, FC 바르셀로나 브랜딩이 특징입니다. "
                "클럽 아이덴티티가 담긴 캐주얼 룩에 적합합니다."
            ),
            "meta_desc": (
                "블라우그라나 FC 바르셀로나 Tee Wall III 반팔 폴로. "
                "폴로 칼라와 클럽 브랜딩."
            ),
        },
    },
    "BLMP813085": {
        "ES": {
            "title": "Camiseta FC Barcelona Tee Wall III Play Antracita",
            "desc": (
                "Camiseta de manga corta de la colección Tee Wall III en antracita. "
                "Gráfico Play en el delantero y detalles de marca FC Barcelona. "
                "Corte regular para un estilo urbano y cómodo."
            ),
            "meta_desc": (
                "Camiseta antracita FC Barcelona Tee Wall III Play de manga corta "
                "con gráfico frontal y corte regular."
            ),
        },
        "CA": {
            "title": "Samarreta FC Barcelona Tee Wall III Play Antracita",
            "desc": (
                "Samarreta de màniga curta de la col·lecció Tee Wall III en antracita. "
                "Gràfic Play al davant i detalls de marca FC Barcelona. "
                "Tall regular per a un estil urbà i còmode."
            ),
            "meta_desc": (
                "Samarreta antracita FC Barcelona Tee Wall III Play de màniga curta "
                "amb gràfic frontal i tall regular."
            ),
        },
        "EN": {
            "title": "FC Barcelona Tee Wall III Play Anthracite T-Shirt",
            "desc": (
                "Short-sleeve t-shirt from the Tee Wall III collection in anthracite. "
                "Play graphic on the front with FC Barcelona branding details. "
                "Regular fit for an easy urban look."
            ),
            "meta_desc": (
                "Anthracite FC Barcelona Tee Wall III Play short-sleeve t-shirt "
                "with front graphic and regular fit."
            ),
        },
        "FR": {
            "title": "T-shirt FC Barcelona Tee Wall III Play Anthracite",
            "desc": (
                "T-shirt à manches courtes de la collection Tee Wall III en anthracite. "
                "Graphique Play sur le devant et détails de marque FC Barcelona. "
                "Coupe regular pour un style urbain confortable."
            ),
            "meta_desc": (
                "T-shirt anthracite FC Barcelona Tee Wall III Play à manches courtes "
                "avec graphique frontal et coupe regular."
            ),
        },
        "DE": {
            "title": "FC Barcelona Tee Wall III Play Anthrazit T-Shirt",
            "desc": (
                "Kurzarm-T-Shirt der Tee Wall III Kollektion in Anthrazit. "
                "Play-Grafik auf der Vorderseite und FC Barcelona Branding-Details. "
                "Regular Fit für einen urbanen Everyday-Look."
            ),
            "meta_desc": (
                "Anthrazitfarbenes FC Barcelona Tee Wall III Play Kurzarm-T-Shirt "
                "mit Frontgrafik und Regular Fit."
            ),
        },
        "JP": {
            "title": "FCバルセロナ Tee Wall III Play アンスラサイト Tシャツ",
            "desc": (
                "Tee Wall IIIコレクションのアンスラサイト半袖Tシャツ。"
                "フロントにPlayグラフィックとFCバルセロナのブランディング。"
                "アーバンなデイリースタイルに合うレギュラーフィット。"
            ),
            "meta_desc": (
                "アンスラサイトのFCバルセロナ Tee Wall III Play半袖Tシャツ。"
                "フロントグラフィックとレギュラーフィット。"
            ),
        },
        "KR": {
            "title": "FC 바르셀로나 Tee Wall III Play 앤트래사이트 티셔츠",
            "desc": (
                "Tee Wall III 컬렉션의 앤트래사이트 반팔 티셔츠입니다. "
                "전면 Play 그래픽과 FC 바르셀로나 브랜딩 디테일이 적용되었습니다. "
                "어반한 데일리 룩에 맞는 레귤러 핏입니다."
            ),
            "meta_desc": (
                "앤트래사이트 FC 바르셀로나 Tee Wall III Play 반팔 티셔츠. "
                "전면 그래픽과 레귤러 핏."
            ),
        },
    },
    "BLMP813086": {
        "ES": {
            "title": "Camiseta FC Barcelona Tee Wall III Pizza Blanca",
            "desc": (
                "Camiseta de manga corta de la colección Tee Wall III en blanco. "
                "Diseño gráfico Pizza en el delantero y marca FC Barcelona. "
                "Corte regular, fácil de combinar."
            ),
            "meta_desc": (
                "Camiseta blanca FC Barcelona Tee Wall III Pizza de manga corta "
                "con gráfico frontal y corte regular."
            ),
        },
        "CA": {
            "title": "Samarreta FC Barcelona Tee Wall III Pizza Blanca",
            "desc": (
                "Samarreta de màniga curta de la col·lecció Tee Wall III en blanc. "
                "Disseny gràfic Pizza al davant i marca FC Barcelona. "
                "Tall regular, fàcil de combinar."
            ),
            "meta_desc": (
                "Samarreta blanca FC Barcelona Tee Wall III Pizza de màniga curta "
                "amb gràfic frontal i tall regular."
            ),
        },
        "EN": {
            "title": "FC Barcelona Tee Wall III Pizza White T-Shirt",
            "desc": (
                "Short-sleeve t-shirt from the Tee Wall III collection in white. "
                "Pizza graphic on the front with FC Barcelona branding. "
                "Regular fit and easy to style."
            ),
            "meta_desc": (
                "White FC Barcelona Tee Wall III Pizza short-sleeve t-shirt "
                "with front graphic and regular fit."
            ),
        },
        "FR": {
            "title": "T-shirt FC Barcelona Tee Wall III Pizza Blanc",
            "desc": (
                "T-shirt à manches courtes de la collection Tee Wall III en blanc. "
                "Graphique Pizza sur le devant et branding FC Barcelona. "
                "Coupe regular, facile à porter au quotidien."
            ),
            "meta_desc": (
                "T-shirt blanc FC Barcelona Tee Wall III Pizza à manches courtes "
                "avec graphique frontal et coupe regular."
            ),
        },
        "DE": {
            "title": "FC Barcelona Tee Wall III Pizza Weiß T-Shirt",
            "desc": (
                "Kurzarm-T-Shirt der Tee Wall III Kollektion in Weiß. "
                "Pizza-Grafik auf der Vorderseite und FC Barcelona Branding. "
                "Regular Fit, leicht zu kombinieren."
            ),
            "meta_desc": (
                "Weißes FC Barcelona Tee Wall III Pizza Kurzarm-T-Shirt "
                "mit Frontgrafik und Regular Fit."
            ),
        },
        "JP": {
            "title": "FCバルセロナ Tee Wall III Pizza ホワイト Tシャツ",
            "desc": (
                "Tee Wall IIIコレクションのホワイト半袖Tシャツ。"
                "フロントにPizzaグラフィックとFCバルセロナのブランディング。"
                "合わせやすいレギュラーフィット。"
            ),
            "meta_desc": (
                "ホワイトのFCバルセロナ Tee Wall III Pizza半袖Tシャツ。"
                "フロントグラフィックとレギュラーフィット。"
            ),
        },
        "KR": {
            "title": "FC 바르셀로나 Tee Wall III Pizza 화이트 티셔츠",
            "desc": (
                "Tee Wall III 컬렉션의 화이트 반팔 티셔츠입니다. "
                "전면 Pizza 그래픽과 FC 바르셀로나 브랜딩이 적용되었습니다. "
                "매치하기 쉬운 레귤러 핏입니다."
            ),
            "meta_desc": (
                "화이트 FC 바르셀로나 Tee Wall III Pizza 반팔 티셔츠. "
                "전면 그래픽과 레귤러 핏."
            ),
        },
    },
    "BLMP813087": {
        "ES": {
            "title": "Camiseta FC Barcelona Tee Wall III Double Sleeve Ecru",
            "desc": (
                "Camiseta de la colección Tee Wall III en ecru con detalle Double Sleeve. "
                "Construcción singular en manga, marca FC Barcelona y corte regular. "
                "Pieza distintiva de la línea Tee Wall."
            ),
            "meta_desc": (
                "Camiseta ecru FC Barcelona Tee Wall III Double Sleeve con "
                "detalle de manga y corte regular."
            ),
        },
        "CA": {
            "title": "Samarreta FC Barcelona Tee Wall III Double Sleeve Ecru",
            "desc": (
                "Samarreta de la col·lecció Tee Wall III en ecru amb detall Double Sleeve. "
                "Construcció singular a la màniga, marca FC Barcelona i tall regular. "
                "Peça distintiva de la línia Tee Wall."
            ),
            "meta_desc": (
                "Samarreta ecru FC Barcelona Tee Wall III Double Sleeve amb "
                "detall de màniga i tall regular."
            ),
        },
        "EN": {
            "title": "FC Barcelona Tee Wall III Double Sleeve Ecru T-Shirt",
            "desc": (
                "Tee Wall III collection t-shirt in ecru with a Double Sleeve detail. "
                "Distinct sleeve construction, FC Barcelona branding and regular fit. "
                "A standout piece from the Tee Wall line."
            ),
            "meta_desc": (
                "Ecru FC Barcelona Tee Wall III Double Sleeve t-shirt with "
                "sleeve detail and regular fit."
            ),
        },
        "FR": {
            "title": "T-shirt FC Barcelona Tee Wall III Double Sleeve Ecru",
            "desc": (
                "T-shirt de la collection Tee Wall III en ecru avec détail Double Sleeve. "
                "Construction de manche singulière, branding FC Barcelona et coupe regular. "
                "Pièce distinctive de la ligne Tee Wall."
            ),
            "meta_desc": (
                "T-shirt ecru FC Barcelona Tee Wall III Double Sleeve avec "
                "détail de manche et coupe regular."
            ),
        },
        "DE": {
            "title": "FC Barcelona Tee Wall III Double Sleeve Ecru T-Shirt",
            "desc": (
                "T-Shirt der Tee Wall III Kollektion in Ecru mit Double-Sleeve-Detail. "
                "Markante Ärmelkonstruktion, FC Barcelona Branding und Regular Fit. "
                "Ein Highlight der Tee Wall Linie."
            ),
            "meta_desc": (
                "Ecru FC Barcelona Tee Wall III Double Sleeve T-Shirt mit "
                "Ärmeldetail und Regular Fit."
            ),
        },
        "JP": {
            "title": "FCバルセロナ Tee Wall III Double Sleeve エクリュ Tシャツ",
            "desc": (
                "Tee Wall IIIコレクションのエクリュTシャツ。Double Sleeveディテール付き。"
                "独自のスリーブ構造、FCバルセロナのブランディング、レギュラーフィット。"
                "Tee Wallラインを象徴する一枚。"
            ),
            "meta_desc": (
                "エクリュのFCバルセロナ Tee Wall III Double Sleeve Tシャツ。"
                "スリーブディテールとレギュラーフィット。"
            ),
        },
        "KR": {
            "title": "FC 바르셀로나 Tee Wall III Double Sleeve 에크루 티셔츠",
            "desc": (
                "Tee Wall III 컬렉션의 에크루 티셔츠로 Double Sleeve 디테일이 특징입니다. "
                "독특한 소매 구조, FC 바르셀로나 브랜딩, 레귤러 핏을 갖췄습니다. "
                "Tee Wall 라인의 시그니처 아이템입니다."
            ),
            "meta_desc": (
                "에크루 FC 바르셀로나 Tee Wall III Double Sleeve 티셔츠. "
                "소매 디테일과 레귤러 핏."
            ),
        },
    },
}


def _is_red(cell) -> bool:
    fill = cell.fill
    if not fill or fill.patternType is None or not fill.fgColor:
        return False
    rgb = getattr(fill.fgColor, "rgb", None)
    return bool(rgb and str(rgb).upper().endswith("FF0000"))


def _set(ws, row: int, letter: str, value: str) -> None:
    ws.cell(row, column_index_from_string(letter)).value = value


def fill_workbook(path: Path, output: Path) -> int:
    wb = load_workbook(path)
    ws = wb.active
    filled = 0

    for row in range(4, ws.max_row + 1):
        ref_cell = ws.cell(row, 1)
        ref = ref_cell.value
        if ref is None:
            continue
        ref_s = str(ref).strip()
        if _is_red(ref_cell):
            continue
        product = PRODUCTS.get(ref_s)
        if not product:
            continue

        for lang, cols in LANG_BLOCKS.items():
            block = product[lang]
            _set(ws, row, cols["title"], block["title"])
            _set(ws, row, cols["desc"], block["desc"])
            _set(ws, row, cols["fit"], FIT[lang])
            if MEASURES_NA[lang]:
                _set(ws, row, cols["measures"], MEASURES_NA[lang])
            _set(ws, row, cols["care"], CARE_TEE[lang])

        for meta_lang, (title_col, desc_col) in META.items():
            src = "JP" if meta_lang == "JA" else "KR" if meta_lang == "KO" else meta_lang
            block = product[src]
            _set(ws, row, title_col, block["title"])
            _set(ws, row, desc_col, block["meta_desc"])

        filled += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return filled


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Source xlsx")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=None,
        help="Output path (default: <input>_filled.xlsx)",
    )
    args = parser.parse_args()
    src = args.input.expanduser().resolve()
    out = args.output.expanduser().resolve() if args.output else src.with_name(f"{src.stem}_filled{src.suffix}")
    n = fill_workbook(src, out)
    print(f"Filled {n} product row(s) → {out}")


if __name__ == "__main__":
    main()
